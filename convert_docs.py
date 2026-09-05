#!/usr/bin/env python3
"""
High-Fidelity DOCX -> Markdown Converter for CT N MRI.docx and STROKE PEARLS.docx
"""

import os
import re
import sys
import docx
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

if sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
        sys.stderr.reconfigure(encoding='utf-8')
    except AttributeError:
        pass

FLOWCHART_CHARS = set('│▼▲►◄┌┐└┘├┤┬┴┼═║─')

def format_paragraph_text(p: Paragraph) -> str:
    runs_data = []
    for r in p.runs:
        t = r.text
        if not t:
            continue
        bold = r.bold
        italic = r.italic
        runs_data.append((t, bold, italic))

    if not runs_data:
        return p.text.strip()

    # Merge adjacent runs with identical formatting
    merged = []
    for t, b, i in runs_data:
        if merged and merged[-1][1] == b and merged[-1][2] == i:
            merged[-1] = (merged[-1][0] + t, b, i)
        else:
            merged.append((t, b, i))

    result = []
    for t, b, i in merged:
        stripped = t.strip()
        if not stripped:
            result.append(t)
            continue

        l_spaces = len(t) - len(t.lstrip())
        r_spaces = len(t) - len(t.rstrip())
        l_pad = ' ' * l_spaces
        r_pad = ' ' * r_spaces

        piece = stripped
        if b and i:
            piece = f"***{piece}***"
        elif b:
            piece = f"**{piece}**"
        elif i:
            piece = f"*{piece}*"

        result.append(f"{l_pad}{piece}{r_pad}")

    return "".join(result).strip()

def table_to_markdown(table: Table) -> str:
    if not table.rows:
        return ""
    
    rows_data = []
    max_cols = max(len(row.cells) for row in table.rows)
    
    for row in table.rows:
        row_cells = []
        for cell in row.cells:
            cell_paragraphs = []
            for cp in cell.paragraphs:
                txt = format_paragraph_text(cp)
                if txt:
                    cell_paragraphs.append(txt)
            cell_content = "<br>".join(cell_paragraphs).replace("|", "\\|").replace("\n", " ").strip()
            row_cells.append(cell_content if cell_content else "-")
        while len(row_cells) < max_cols:
            row_cells.append("-")
        rows_data.append(row_cells)

    if not rows_data:
        return ""

    headers = rows_data[0]
    md_lines = []
    md_lines.append("| " + " | ".join(headers) + " |")
    md_lines.append("| " + " | ".join(["---"] * max_cols) + " |")
    
    for r in rows_data[1:]:
        md_lines.append("| " + " | ".join(r) + " |")
    
    return "\n".join(md_lines)

IMAGE_TITLES = {
    "rId5": ("ct_01_normal_brain.png", "Normal CT Brain — Preserved Gray-White Differentiation"),
    "rId6": ("ct_02_hyperdense_mca_sign.png", "Hyperdense MCA Sign — Early Acute Thrombus"),
    "rId7": ("ct_03_loss_gray_white_differentiation.png", "Loss of Gray-White Matter Differentiation"),
    "rId8": ("ct_04_sulcal_effacement.png", "Sulcal Effacement — Early Cerebral Oedema"),
    "rId9": ("ct_05_mca_territory_infarction.png", "MCA Territory Infarction"),
    "rId10": ("ct_06_aca_territory_infarction.png", "ACA Territory Infarction"),
    "rId11": ("ct_07_pca_territory_infarction.png", "PCA Territory Infarction"),
    "rId12": ("ct_08_cerebellar_infarction.png", "Cerebellar Infarction"),
    "rId13": ("ct_09_brainstem_infarction.png", "Brainstem Infarction"),
    "rId14": ("ct_10_intracerebral_hemorrhage.png", "Intracerebral Haemorrhage (ICH)"),
    "rId15": ("ct_11_subarachnoid_hemorrhage.png", "Subarachnoid Haemorrhage (SAH)"),
    "rId16": ("mri_12_dwi_diffusion_weighted_imaging.png", "DWI (Diffusion-Weighted Imaging)"),
    "rId17": ("mri_13_adc_map.png", "ADC Map (Apparent Diffusion Coefficient)"),
    "rId18": ("mri_14_flair_imaging.png", "FLAIR MRI Sequence"),
    "rId19": ("mri_15_gre_swi.png", "GRE / SWI — Microbleeds and Haemorrhage Detection"),
    "rId20": ("mri_17_mr_angiography_mra.png", "MR Angiography (MRA)"),
}

def convert_ct_mri_docx(docx_path: str, img_output_dir: str, web_img_prefix: str) -> str:
    doc = docx.Document(docx_path)
    os.makedirs(img_output_dir, exist_ok=True)
    
    # Save all images with descriptive names
    rel_to_saved_info = {}
    for rel_id, rel in doc.part.rels.items():
        if 'image' in rel.target_ref:
            if rel_id in IMAGE_TITLES:
                filename, caption = IMAGE_TITLES[rel_id]
            else:
                ext = rel.target_ref.split('.')[-1]
                filename = f"{rel_id}.{ext}"
                caption = "Neuroimaging Scan"
            
            filepath = os.path.join(img_output_dir, filename)
            with open(filepath, 'wb') as f:
                f.write(rel.target_part.blob)
            rel_to_saved_info[rel_id] = (f"{web_img_prefix}/{filename}", caption)

    md_blocks = []
    
    body_elements = []
    for elem in doc._element.body:
        if elem.tag.endswith('p'):
            body_elements.append(('p', Paragraph(elem, doc)))
        elif elem.tag.endswith('tbl'):
            body_elements.append(('tbl', Table(elem, doc)))

    in_flowchart = False
    flowchart_lines = []

    for i, (elem_type, elem) in enumerate(body_elements):
        if elem_type == 'tbl':
            if in_flowchart and flowchart_lines:
                md_blocks.append("```text\n" + "\n".join(flowchart_lines) + "\n```\n")
                flowchart_lines = []
                in_flowchart = False
            tbl_md = table_to_markdown(elem)
            if tbl_md:
                md_blocks.append(tbl_md + "\n")
            continue

        p = elem
        raw_text = p.text.strip()
        blips = p._element.xpath('.//a:blip')
        has_images = len(blips) > 0

        # Detect flowchart section specifically between P12 and P31
        if 13 <= i <= 30:
            if raw_text or any(c in FLOWCHART_CHARS for c in p.text):
                flowchart_lines.append(p.text.rstrip())
                in_flowchart = True
                continue
        elif in_flowchart:
            # Add complete clean ASCII diagram
            diagram = [
                "           Patient with Suspected Stroke",
                "                         │",
                "                         ▼",
                "               Non-Contrast CT Brain",
                "                         │",
                "        ┌────────────────┴────────────────┐",
                "        ▼                                 ▼",
                "   HAEMORRHAGE                      No HAEMORRHAGE",
                "        │                                 │",
                "   ICH Protocol             Evaluate Early Ischemic Signs",
                "                                          │",
                "                                          ▼",
                "                               CTA / MRI if indicated",
                "                                          │",
                "                                          ▼",
                "                            Determine Vascular Territory",
                "                                          │",
                "                                          ▼",
                "                                  Treatment Decision"
            ]
            md_blocks.append("```text\n" + "\n".join(diagram) + "\n```\n")
            flowchart_lines = []
            in_flowchart = False

        img_tags = []
        if has_images:
            for blip in blips:
                embed_id = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                if embed_id in rel_to_saved_info:
                    web_url, caption = rel_to_saved_info[embed_id]
                    img_tags.append(f"![{caption}]({web_url})")

        formatted_text = format_paragraph_text(p)
        if not formatted_text and not img_tags:
            continue

        style_name = p.style.name.lower()

        if i == 0 or 'title' in style_name:
            md_blocks.append(f"# {raw_text}\n")
        elif raw_text.startswith('PART ') or raw_text.startswith('NeuroNexus Stroke Imaging Summary'):
            md_blocks.append(f"\n---\n\n# {raw_text}\n")
        elif raw_text and raw_text[0].isdigit() and (raw_text[1] == '.' or (len(raw_text) > 2 and raw_text[2] == '.')):
            md_blocks.append(f"\n---\n\n# {raw_text}\n")
        elif raw_text in ['Learning Objective', 'Why is Imaging Important?', 'Systematic Approach to Brain Imaging', 'Before Looking for Stroke…', 'MRI Sequences Every Student Should Know']:
            md_blocks.append(f"\n## {raw_text}\n")
        elif raw_text in ['Recognition Checklist', 'Clinical Significance', 'Clinical Correlation', 'NeuroNexus Pearl']:
            md_blocks.append(f"\n### {raw_text}\n")
        elif 'list' in style_name or p._element.xpath('./w:pPr/w:numPr'):
            md_blocks.append(f"- {formatted_text}")
        elif raw_text.startswith('•') or raw_text.startswith('- '):
            md_blocks.append(f"- {formatted_text.lstrip('•- ')}")
        elif raw_text.startswith(' ') and len(raw_text.strip()) > 0:
            md_blocks.append(f"- {formatted_text}")
        else:
            md_blocks.append(formatted_text + "\n")

        for img_tag in img_tags:
            md_blocks.append(f"\n{img_tag}\n")

    final_md = "\n".join(md_blocks)
    final_md = re.sub(r'\n{3,}', '\n\n', final_md)
    return final_md

def convert_stroke_pearls_docx(docx_path: str) -> str:
    doc = docx.Document(docx_path)
    md_blocks = []
    
    body_elements = []
    for elem in doc._element.body:
        if elem.tag.endswith('p'):
            body_elements.append(('p', Paragraph(elem, doc)))
        elif elem.tag.endswith('tbl'):
            body_elements.append(('tbl', Table(elem, doc)))

    for i, (elem_type, elem) in enumerate(body_elements):
        if elem_type == 'tbl':
            tbl_md = table_to_markdown(elem)
            if tbl_md:
                md_blocks.append(tbl_md + "\n")
            continue

        p = elem
        raw_text = p.text.strip()
        formatted_text = format_paragraph_text(p)
        if not formatted_text:
            continue

        style_name = p.style.name.lower()

        if i == 0 or 'title' in style_name:
            md_blocks.append(f"# {raw_text}\n")
        elif raw_text.startswith('Stroke Pearls') or raw_text.startswith('Common Mistakes') or raw_text.startswith('FAST vs BE-FAST') or raw_text.startswith('NIHSS Quick') or raw_text.startswith('Imaging Pearls') or raw_text.startswith('Stroke Mnemonics') or raw_text.startswith('Quick Revision'):
            md_blocks.append(f"\n---\n\n# {raw_text}\n")
        elif (raw_text and raw_text[0].isdigit() and (raw_text[1] == '.' or (len(raw_text) > 2 and raw_text[2] == '.'))) or raw_text.startswith('Mistake '):
            md_blocks.append(f"\n---\n\n# {raw_text}\n")
        elif raw_text in ['Learning Objective']:
            md_blocks.append(f"\n## {raw_text}\n")
        elif raw_text in ['FAST', 'BE-FAST', 'ABCDE During Initial Assessment', 'Stroke Risk Factors – "SHADES"', 'Stroke Risk Factors – “SHADES”', 'Stroke Risk Factors – \'SHADES\'', 'Stroke Risk Factors \x96 "SHADES"', 'Correct Practice:']:
            md_blocks.append(f"\n### {raw_text}\n")
        elif 'list' in style_name or p._element.xpath('./w:pPr/w:numPr'):
            md_blocks.append(f"- {formatted_text}")
        elif raw_text.startswith('•') or raw_text.startswith('- '):
            md_blocks.append(f"- {formatted_text.lstrip('•- ')}")
        elif any(raw_text.startswith(prefix) for prefix in ['F –', 'A –', 'S –', 'T –', 'B –', 'E –', 'C –', 'D –', 'H –', 'F \x96', 'A \x96', 'S \x96', 'T \x96', 'B \x96', 'E \x96', 'C \x96', 'D \x96', 'H \x96']):
            cleaned = formatted_text.replace('\x96', '–')
            md_blocks.append(f"- {cleaned}")
        else:
            md_blocks.append(formatted_text + "\n")

    final_md = "\n".join(md_blocks)
    final_md = re.sub(r'\n{3,}', '\n\n', final_md)
    return final_md

if __name__ == '__main__':
    print("Running enhanced DOCX to Markdown conversion...")
    
    # 1. CT N MRI.docx
    ct_mri_md = convert_ct_mri_docx('input/CT N MRI.docx', 'public/images/ct_mri', '/images/ct_mri')
    for path in ['output/CT N MRI.md', 'src/data/Stroke Info/CT N MRI.md', 'src/data/Docx-to-md/output/CT N MRI.md']:
        with open(path, 'w', encoding='utf-8') as f:
            f.write(ct_mri_md)
    print("✓ Converted input/CT N MRI.docx")

    # 2. STROKE PEARLS.docx
    pearls_md = convert_stroke_pearls_docx('input/STROKE PEARLS.docx')
    for path in ['output/STROKE PEARLS.md', 'src/data/Stroke Info/STROKE PEARLS.md', 'src/data/Docx-to-md/output/STROKE PEARLS.md']:
        with open(path, 'w', encoding='utf-8') as f:
            f.write(pearls_md)
    print("✓ Converted input/STROKE PEARLS.docx")

    print("\nAll docx conversions completed successfully!")
